"""Applies the owner's filled review sheet (the 128 options that were waiting for a price).

    python -I catalogue/build_owner_review_prices.py "catalogue/audit/128_OPTIONS_MRP_FILLED (1).xlsx"

Each row is matched to its variant through the original Excel reference (workbook, sheet, row) and must agree
with the source on product name, option and original selling price. Rules (owner's decision 11-Oct-2026):
- MRP = "My Corrected MRP", kept exactly.
- Selling price = "My Approved Selling Price" when filled, otherwise 95% of the corrected MRP (2 decimals, half up).
- Rejected and reported (never guessed): no or invalid MRP/price, price above MRP, more than 2 decimals, a row
  that doesn't match its source row, the same option with different values on two sheets.

Writes staging/owner_review_prices.csv (import_catalog sheet; absolute values, so re-running changes nothing)
and audit/owner_review_result.xlsx.
"""

import csv
import sys
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

sys.stdout.reconfigure(encoding="utf-8")
HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
import build_full_catalogue as b  # noqa: E402

SHEETS = ("121 Missing MRP", "7 Price Conflicts", "All 128 Options")


def money(value):
    if value in (None, ""):
        return None
    try:
        d = Decimal(str(value).replace(",", "").replace("₹", "").strip())
    except InvalidOperation:
        raise ValueError(f"not a number: {value!r}") from None
    if d <= 0:
        raise ValueError(f"must be above zero: {value}")
    if d != d.quantize(Decimal("0.01")):
        raise ValueError(f"more than 2 decimals: {value}")
    return d.quantize(Decimal("0.01"))


def main(path: Path):
    review = load_workbook(path, data_only=True)
    products, _ = b.load_products()
    source = {}  # (workbook file, sheet, row) → (product, option)
    for p in products:
        workbook, sheet = p["var_sheet"].split(" › ")
        for v in p["variants"]:
            source[(workbook.rsplit("/", 1)[-1], sheet, v["_row"])] = (p, v)
    staged = {}
    for name in ("full_catalogue_products.csv", "full_catalogue_pilot_prices.csv"):
        for r in csv.DictReader((HERE / "staging" / name).open(encoding="utf-8")):
            staged[r["variant_sku"]] = r
    report = list(load_workbook(HERE / "audit/import_report.xlsx", read_only=True)["Options"].iter_rows(values_only=True))
    head = report[0]
    sku_of = {(r[head.index("product")], r[head.index("option")]): r[head.index("variant_sku")] for r in report[1:]}

    seen, errors, results = {}, [], []
    for sheet_name in SHEETS:
        ws = review[sheet_name]
        rows = list(ws.iter_rows(values_only=True))
        cols = rows[0]
        for n, values in enumerate(rows[1:], start=2):
            if not any(v not in (None, "") for v in values):
                continue
            d = dict(zip(cols, values))
            where = f"{sheet_name} row {n}"
            ref = (d["Excel Workbook Name"], d["Excel Sheet Name"], d["Excel Row Number"])
            try:
                if ref not in source:
                    raise ValueError(f"no source row {ref}")
                p, v = source[ref]
                option = b.option_label(v)[0]
                if (d["Product Name"], d["Variant/Option Name"]) != (p["name"], option):
                    raise ValueError(f"name/option differ from the source row ({p['name']} / {option})")
                if d["Original Selling Price"] != v.get("Selling Price (₹)"):
                    raise ValueError(f"original price {d['Original Selling Price']} differs from source {v.get('Selling Price (₹)')}")
                mrp = money(d["My Corrected MRP"])
                if mrp is None:
                    raise ValueError("My Corrected MRP is empty")
                price = money(d["My Approved Selling Price"])
                rule = "price as filled in"
                if price is None:
                    price, rule = (mrp * Decimal("0.95")).quantize(Decimal("0.01"), ROUND_HALF_UP), "95% of MRP"
                if price > mrp:
                    raise ValueError(f"price {price} is above MRP {mrp}")
                sku = sku_of[(p["name"], option)]
            except (ValueError, KeyError) as exc:
                errors.append({"where": where, "product": d.get("Product Name"), "option": d.get("Variant/Option Name"), "error": str(exc)})
                continue
            key = sku
            if key in seen:
                if seen[key] != (mrp, price):
                    errors.append({"where": where, "product": p["name"], "option": option,
                                   "error": f"differs from {seen[key]} on another sheet"})  # fmt: skip
                continue
            seen[key] = (mrp, price)
            results.append({"variant_sku": sku, "product_sku": staged[sku]["product_sku"], "product": p["name"],
                            "option": option, "problem": d["Problem Type"], "excel price": v.get("Selling Price (₹)"),
                            "excel MRP": v.get("MRP (₹)"), "new MRP": mrp, "new price": price, "price rule": rule,
                            "discount %": float(((mrp - price) * 100 / mrp).quantize(Decimal("0.1"))),
                            "MRP above Excel MRP": "yes" if isinstance(v.get("MRP (₹)"), (int, float)) and mrp > Decimal(str(v["MRP (₹)"])) else ""})  # fmt: skip
    # a variant must not be applied if any of its rows was rejected
    bad = {sku_of.get((e["product"], e["option"])) for e in errors}
    results = [r for r in results if r["variant_sku"] not in bad]
    with (HERE / "staging/owner_review_prices.csv").open("w", encoding="utf-8", newline="") as fh:
        w = csv.DictWriter(fh, ["product_sku", "variant_sku", "variant_name", "attributes", "mrp", "price", "price_approved", "price_source"])
        w.writeheader()
        for r in results:
            s = staged[r["variant_sku"]]
            w.writerow({"product_sku": r["product_sku"], "variant_sku": r["variant_sku"], "variant_name": s["variant_name"],
                        "attributes": s["attributes"], "mrp": str(r["new MRP"]), "price": str(r["new price"]),
                        "price_approved": "yes",
                        "price_source": f"Owner review sheet {path.name}, 11-Oct-2026 (MRP from owner; {r['price rule']})"})  # fmt: skip
    out = Workbook()
    for title, data in (("Applied", results), ("Errors", errors)):
        ws = out.create_sheet(title)
        if data:
            ws.append(list(data[0]))
            for c in ws[1]:
                c.font = Font(bold=True)
            for r in data:
                ws.append([float(x) if isinstance(x, Decimal) else x for x in r.values()])
    out.remove(out["Sheet"])
    out.save(HERE / "audit/owner_review_result.xlsx")
    print(f"unique options {len(results)} valid, errors {len(errors)}; "
          f"missing-MRP {sum(r['problem'] == 'MRP REQUIRED' for r in results)}, conflicts {sum(r['problem'] == 'PRICE/MRP CONFLICT' for r in results)}; "
          f"rules {dict((k, sum(r['price rule'] == k for r in results)) for k in {r['price rule'] for r in results})}; "
          f"MRP raised above Excel MRP: {sum(bool(r['MRP above Excel MRP']) for r in results)}")
    for e in errors:
        print("ERROR", e)


if __name__ == "__main__":
    main(Path(sys.argv[1]))
