"""Owner's pricing decision (11-Oct-2026), applied to the full catalogue staging data.

    python -I catalogue/build_price_approval.py     # after build_full_catalogue.py

Decision: the Excel ("Full Price List") prices are approved for the shop; the Excel MRP stays the MRP and the
selling price is 95% of the Excel selling price, rounded to 2 decimals (half up).

Writes staging/full_catalogue_approved_prices.csv (import_catalog sheet, price columns only, so descriptions
and other product data are untouched) and audit/price_approval_report.xlsx. Per option:
- Excel price and MRP present, 95% price <= MRP: approved (existing variant updated, or created if it was
  waiting because the full Excel price was above the MRP).
- Excel MRP missing: left waiting (an MRP is not invented).
- 95% price still above the Excel MRP: left waiting (the source data conflicts).
"""

import csv
import sys
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

sys.stdout.reconfigure(encoding="utf-8")
HERE = Path(__file__).resolve().parent
STAGING, AUDIT = HERE / "staging", HERE / "audit"
SHARE = Decimal("0.95")
SOURCE = "Excel Full Price List (HomeRun 10-Oct-2026); approved by owner 11-Oct-2026; price = 95% of Excel price"


def dec(v):
    return None if v in (None, "") else Decimal(str(v))


def main():
    wb = load_workbook(AUDIT / "import_report.xlsx", read_only=True)
    rows = list(wb["Options"].iter_rows(values_only=True))
    head = rows[0]
    options = {r[head.index("variant_sku")]: dict(zip(head, r)) for r in rows[1:] if r[head.index("variant_sku")]}
    staged = []
    for name in ("full_catalogue_products.csv", "full_catalogue_pilot_prices.csv"):
        staged += list(csv.DictReader((STAGING / name).open(encoding="utf-8")))
    out, report = [], []
    for row in staged:
        o = options[row["variant_sku"]]
        price, mrp = dec(o["HomeRun price"]), dec(o["HomeRun MRP"])
        new_price = (price * SHARE).quantize(Decimal("0.01"), ROUND_HALF_UP) if price else None
        rec = {"product": o["product"], "product_sku": row["product_sku"], "variant_sku": row["variant_sku"],
               "option": o["option"], "Excel price": price, "Excel MRP": mrp, "selling price (95%)": new_price,
               "discount vs MRP %": None, "result": "", "note": ""}  # fmt: skip
        if price is None:
            rec.update(result="waiting", note="no selling price in Excel")
        elif mrp is None:
            rec.update(result="waiting", note="no MRP in Excel; MRP not invented")
        elif new_price > mrp:
            rec.update(result="waiting", note=f"95% price ₹{new_price} is still above the Excel MRP ₹{mrp}")
        else:
            rec.update(
                result="approved",
                note="was above MRP at full Excel price; valid after the 5% reduction" if price > mrp else "",
                **{"discount vs MRP %": float(((mrp - new_price) * 100 / mrp).quantize(Decimal("0.1")))},
            )
            out.append({
                "product_sku": row["product_sku"], "variant_sku": row["variant_sku"],
                "variant_name": row["variant_name"], "attributes": row["attributes"],
                "mrp": str(mrp), "price": str(new_price), "price_approved": "yes", "price_source": SOURCE,
            })  # fmt: skip
        report.append(rec)
    with (STAGING / "full_catalogue_approved_prices.csv").open("w", encoding="utf-8", newline="") as fh:
        w = csv.DictWriter(fh, list(out[0]))
        w.writeheader()
        w.writerows(out)
    rb = Workbook()
    ws = rb.active
    ws.title = "Options"
    cols = list(report[0])
    ws.append(cols)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in report:
        ws.append([float(v) if isinstance(v, Decimal) else v for v in (r[c] for c in cols)])
    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = "A2"
    rb.save(AUDIT / "price_approval_report.xlsx")
    count = lambda res: sum(1 for r in report if r["result"] == res)  # noqa: E731
    print(f"options {len(report)}: approved {count('approved')}, waiting {count('waiting')} "
          f"(no MRP {sum(1 for r in report if 'no MRP' in r['note'])}, "
          f"still above MRP {sum(1 for r in report if 'still above' in r['note'])}); "
          f"rescued from above-MRP {sum(1 for r in report if r['note'].startswith('was above'))}")


if __name__ == "__main__":
    main()
