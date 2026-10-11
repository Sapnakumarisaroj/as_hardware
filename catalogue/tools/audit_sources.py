"""Stage 1 read-only audit of assets/as-hardware-images: workbooks, price lists, images.

usage: python -I audit.py <assets root> <out.xlsx>
"""
import difflib
import hashlib
import json
import os
import re
import sys
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

sys.stdout.reconfigure(encoding="utf-8")
root, out = Path(sys.argv[1]), Path(sys.argv[2])


def norm(s):
    s = str(s or "").lower().replace("&", " and ").replace("°", "")
    s = re.sub(r"[^a-z0-9]+", " ", s)
    return " ".join(s.split())


def money(v):
    if v in (None, "", "—", "-", "Not listed"):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    m = re.search(r"[\d,]+(?:\.\d+)?", str(v))
    return float(m.group().replace(",", "")) if m else None


def header_index(rows):
    for i, r in enumerate(rows[:6]):
        if r and "Product Name" in [str(c).strip() if c else "" for c in r]:
            return i
    return None


def table(ws):
    rows = list(ws.iter_rows(values_only=True))
    h = header_index(rows)
    if h is None:
        return [], []
    head = [str(c).strip() if c else f"col{j}" for j, c in enumerate(rows[h])]
    recs = []
    for i, r in enumerate(rows[h + 1:], start=h + 2):
        d = dict(zip(head, r))
        if not d.get("Product Name"):
            continue
        d["_row"] = i
        recs.append(d)
    return head, recs


PRICE_LINE = re.compile(r"^(?P<label>.*?):\s*₹\s?(?P<price>[\d,]+(?:\.\d+)?)(?:\s*\(MRP\s*₹\s?(?P<mrp>[\d,]+(?:\.\d+)?)\))?")


def parse_full_list(text):
    out = []
    for line in str(text or "").splitlines():
        line = line.strip()
        if not line:
            continue
        if "→" in line:
            thick, rest = line.split("→", 1)
            for part in rest.split("|"):
                m = PRICE_LINE.match(part.strip())
                if m:
                    out.append((f"{thick.strip()} / {m['label'].strip()}", money(m["price"]), money(m["mrp"]) if m["mrp"] else None, line))
                else:
                    out.append((None, None, None, part))
            continue
        m = PRICE_LINE.match(line)
        if m:
            out.append((m["label"].strip(), money(m["price"]), money(m["mrp"]) if m["mrp"] else None, line))
        else:
            out.append((None, None, None, line))
    return out


products, variants, issues, images = [], [], [], []
files_summary = []
image_files = sorted(p for p in root.rglob("*.webp"))
img_hash = {p: hashlib.sha256(p.read_bytes()).hexdigest() for p in image_files}


def issue(kind, severity, where, product, detail):
    issues.append({"severity": severity, "kind": kind, "file": where, "product": product, "detail": detail})


for f in sorted(root.rglob("*.xlsx")):
    rel = f.relative_to(root).as_posix()
    folder = f.parent
    top, sub = f.relative_to(root).parts[0], f.relative_to(root).parts[1]
    wb = load_workbook(f, read_only=True, data_only=True)
    _, prod_recs = table(wb["Products"])
    var_sheet = next(ws for ws in wb.worksheets if ws.title != "Products" and "SKU" in (table(ws)[0] or []))
    vhead, var_recs = table(var_sheet)
    price_text_col = next(c for c in (prod_recs[0] if prod_recs else {}) if str(c).startswith("Full Price List"))
    by_name = defaultdict(list)
    for v in var_recs:
        by_name[v["Product Name"].strip()].append(v)
    files_summary.append({"file": rel, "products": len(prod_recs), "variant rows": len(var_recs),
                          "price sheet": var_sheet.title, "images in folder": len(list(folder.glob("*.webp")))})
    names_seen = Counter(p["Product Name"].strip() for p in prod_recs)
    for n, c in names_seen.items():
        if c > 1:
            issue("duplicate product row", "high", rel, n, f"listed {c} times on Products sheet")
    for p in prod_recs:
        name = p["Product Name"].strip()
        pid = f"{Path(rel).stem}#{p.get('#')}"
        vs = by_name.get(name, [])
        listed = parse_full_list(p.get(price_text_col))
        unparsed = [l for l in listed if l[0] is None]
        n_opts = p.get("No. of Options") or p.get("No. of Pack Sizes") or p.get("No. of Variants")
        rec = {
            "id": pid, "file": rel, "category": top, "subcategory folder": sub, "name": name,
            "brand": p.get("Brand"), "product type": p.get("Product Type"), "description (HomeRun)": p.get("Description"),
            "options listed": n_opts, "variant rows": len(vs), "price lines": len(listed),
        }
        if not vs:
            issue("no price rows", "high", rel, name, "product has no rows on the price sheet")
        if n_opts and isinstance(n_opts, (int, float)) and int(n_opts) != len(vs):
            issue("option count mismatch", "medium", rel, name, f"Products says {n_opts} options, price sheet has {len(vs)}")
        for l in unparsed:
            if "₹" in l[3]:
                issue("price text not parsed", "low", rel, name, l[3][:120])
        # cross-check every variant row against the 'Full Price List' text
        text_prices = [(norm(l[0]), l[1], l[2]) for l in listed if l[0]]
        for v in vs:
            label = v.get("Size / Variety") or v.get("Pack Size") or v.get("Thickness / Variety")
            if "Size" in v and v.get("Thickness / Variety"):
                label = f"{v.get('Thickness / Variety')} / {v.get('Size')}"
            price, mrp = money(v.get("Selling Price (₹)")), money(v.get("MRP (₹)"))
            stock = v.get("In Stock")
            match = [t for t in text_prices if t[0] == norm(label)]
            if not match and len(text_prices) == 1 and len(vs) == 1:
                match = text_prices
            if not match:
                # plywood lists thickness lines that may contain both sizes; try containment
                match = [t for t in text_prices if norm(label).split(" / ")[0] in t[0] or t[0] in norm(label)]
            status = "ok"
            if price is None:
                status = "missing price"
                issue("missing selling price", "high", rel, name, f"{label}: no selling price")
            if mrp is None:
                issue("MRP not listed", "medium", rel, name, f"{label}: HomeRun shows no MRP (price ₹{price:g})" if price else f"{label}: no MRP")
            if price and mrp and price > mrp:
                status = "price above MRP"
                issue("price above MRP", "high", rel, name, f"{label}: ₹{price:g} > MRP ₹{mrp:g}")
            exact = [t for t in match if t[1] == price and (t[2] == mrp or t[2] is None and mrp is None)]
            if text_prices and not exact:
                if match:
                    status = "conflict"
                    issue("price conflict (Products text vs price sheet)", "high", rel, name,
                          f"{label}: sheet ₹{price} / MRP {mrp}; text says " + "; ".join(f"₹{t[1]:g}" + (f" (MRP ₹{t[2]:g})" if t[2] else "") for t in match[:3]))
                else:
                    status = "not in price text"
            variants.append({
                "product id": pid, "file": rel, "product": name, "brand": p.get("Brand"), "option": label,
                "HomeRun SKU": v.get("SKU"), "selling price": price, "MRP": mrp,
                "bulk offer": v.get("Bulk Price Offer"), "in stock (HomeRun)": stock, "check": status,
            })
        products.append(rec)
    # images in this folder
    folder_imgs = sorted(folder.glob("*.webp"))
    names = {norm(p["Product Name"]): p["Product Name"].strip() for p in prod_recs}
    for img in folder_imgs:
        stem = norm(img.stem)
        base = re.sub(r"( \d{1,2}| pic \d+.*| copy)$", "", stem)
        best, score, how = None, 0.0, ""
        if stem in names:
            best, score, how = names[stem], 1.0, "exact name"
        elif base in names:
            best, score, how = names[base], 0.99, "exact name + number (gallery)"
        else:
            for k, n in names.items():
                r = difflib.SequenceMatcher(None, base, k).ratio()
                toks_img, toks_p = set(base.split()), set(k.split())
                cov = len(toks_img & toks_p) / max(1, len(toks_img))
                s = max(r, cov * 0.97 if toks_img <= toks_p or cov >= 0.9 else r)
                if s > score:
                    best, score, how = n, s, "fuzzy"
        verdict = "match" if score >= 0.99 else "likely — confirm" if score >= 0.85 else "uncertain" if score >= 0.6 else "no match"
        images.append({"folder": folder.relative_to(root).as_posix(), "image file": img.name, "product": best if score >= 0.6 else "",
                       "score": round(score, 2), "method": how, "verdict": verdict, "sha256": img_hash[img][:12]})
    for p in products[-len(prod_recs):]:
        p["images matched"] = sum(1 for i in images if i["product"] == p["name"] and i["verdict"] == "match" and i["folder"] == folder.relative_to(root).as_posix())
        p["images likely"] = sum(1 for i in images if i["product"] == p["name"] and i["verdict"] == "likely — confirm" and i["folder"] == folder.relative_to(root).as_posix())
        if p["images matched"] == 0 and p["images likely"] == 0:
            issue("no image", "medium", rel, p["name"], "no image file matched with confidence")

all_names = {norm(p["name"]): p for p in products}
for i in images:
    if i["verdict"] in ("no match", "uncertain"):
        base = re.sub(r"( \d{1,2}| pic \d+.*| copy)$", "", norm(i["image file"].rsplit(".", 1)[0]))
        stem = norm(i["image file"].rsplit(".", 1)[0])
        if stem in all_names:
            base = stem
        if base in all_names:
            i.update(product=all_names[base]["name"], score=1.0, method=f"exact name, product is in {all_names[base]['file'].split('/')[1]}", verdict="match")
        elif i["verdict"] == "no match":
            i["method"] = "no product with this name in any workbook (no price data)"
for p in products:
    p["images matched"] = sum(1 for i in images if i["product"] == p["name"] and i["verdict"] == "match")
    p["images likely"] = sum(1 for i in images if i["product"] == p["name"] and i["verdict"] == "likely — confirm")
issues[:] = [x for x in issues if x["kind"] != "no image"]
for p in products:
    if not p["images matched"] and not p["images likely"]:
        issue("no image", "medium", p["file"], p["name"], "no image file matched with confidence")
# cross-file checks
name_files = defaultdict(set)
for p in products:
    name_files[norm(p["name"])].add(p["file"])
for n, fs in name_files.items():
    if len(fs) > 1:
        issue("same product in several files", "high", "; ".join(sorted(fs)), n, "listed in more than one category file")
sku_count = Counter(v["HomeRun SKU"] for v in variants if v["HomeRun SKU"])
for s, c in sku_count.items():
    if c > 1:
        issue("repeated HomeRun SKU", "high", "", s, f"SKU appears on {c} variant rows: " + "; ".join(f"{v['product']} / {v['option']}" for v in variants if v["HomeRun SKU"] == s)[:250])
hash_groups = defaultdict(list)
for i in images:
    hash_groups[i["sha256"]].append(i)
for h, grp in hash_groups.items():
    if len(grp) > 1:
        issue("identical image files", "medium", "; ".join(g["folder"] for g in grp), grp[0]["image file"], " = ".join(g["image file"] for g in grp))
        prods = {g["product"] for g in grp if g["product"]}
        if len(prods) > 1:
            issue("same picture on different products", "high", "", h, " / ".join(prods))
for v in variants:
    pass
# product-level image conflicts: two different products claim same image handled above

# workbook
wb = Workbook()
ws = wb.active
ws.title = "Summary"
sev = Counter(i["severity"] for i in issues)
kinds = Counter((i["severity"], i["kind"]) for i in issues)
summary_rows = [
    ["Stage 1 audit — AS Hardware catalogue sources (read-only)"],
    ["Source of every workbook", "HomeRun (home-run.co) export, data captured 10-Oct-2026 — another retailer's listings and prices"],
    [],
    ["Workbook", "Products", "Variant rows", "Price sheet", "Images in folder"],
    *[[s["file"], s["products"], s["variant rows"], s["price sheet"], s["images in folder"]] for s in files_summary],
    ["TOTAL", sum(s["products"] for s in files_summary), sum(s["variant rows"] for s in files_summary), "", len(image_files)],
    [],
    ["Issue", "Severity", "Count"],
    *[[k, s, c] for (s, k), c in sorted(kinds.items(), key=lambda x: ({"high": 0, "medium": 1, "low": 2}[x[0][0]], -x[1]))],
    [],
    ["Images", "Count"],
    *[[k, c] for k, c in Counter(i["verdict"] for i in images).items()],
]
for r in summary_rows:
    ws.append(r)
ws["A1"].font = Font(bold=True, size=14)


def sheet(title, recs):
    s = wb.create_sheet(title)
    if not recs:
        return
    cols = list(recs[0].keys())
    s.append(cols)
    for c in s[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DDDDDD")
    for r in recs:
        s.append([r.get(c) for c in cols])
    for j, c in enumerate(cols, 1):
        width = min(60, max(10, max(len(str(r.get(c) or "")) for r in recs[:300]) + 2, len(c) + 2))
        s.column_dimensions[get_column_letter(j)].width = width
    s.freeze_panes = "A2"
    s.auto_filter.ref = s.dimensions


sheet("Issues", sorted(issues, key=lambda i: ({"high": 0, "medium": 1, "low": 2}[i["severity"]], i["kind"])))
sheet("Products", products)
sheet("Variants & prices", variants)
sheet("Images", images)
out.parent.mkdir(parents=True, exist_ok=True)
wb.save(out)

print(json.dumps({"files": files_summary, "issues": {f"{s}: {k}": c for (s, k), c in kinds.items()},
                  "images": Counter(i["verdict"] for i in images),
                  "products": len(products), "variants": len(variants),
                  "variant checks": Counter(v["check"] for v in variants),
                  "no-MRP variants": sum(1 for v in variants if v["MRP"] is None),
                  "products w/o confident image": sum(1 for p in products if not p["images matched"] and not p["images likely"])},
                 indent=1, ensure_ascii=False, default=str))
