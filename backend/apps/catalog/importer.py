"""Catalogue import from a spreadsheet (CSV or Excel .xlsx), one row per variant.

    python manage.py import_catalog products.xlsx            # dry run: check every row, save nothing
    python manage.py import_catalog products.xlsx --commit   # save, all or nothing
    POST /api/admin/catalog/import/                          # the same from the dashboard

The column reference is `COLUMNS` below (and docs/CATALOG_IMPORT.md). Rules:
- Rows with the same `product_sku` are the variants of one product. Product columns (name, category,
  brand, …) may be left blank on the product's later rows, but where they are filled they must match.
- Upsert by SKU: an existing product/variant SKU is updated, a new one is created. A column that is not
  in the file is never touched; a column that is in the file is applied as written (a blank optional
  cell clears the value). Exceptions: a blank `stock` cell changes no stock, a blank `is_published`,
  `unit` or `max_order_quantity` keeps the current value (new products: unpublished, piece, 50).
- Nothing is invented: a new product needs its name, approved category and GST rate; a new variant
  needs its MRP and price.
- Every row is validated before anything is saved and errors are reported per row and column. The
  import then runs in one transaction: any failure rolls everything back. A dry run (the default)
  runs the whole import and rolls it back, so it reports exactly what a real import would do.
- Stock (`BLR-01:25`) is the wanted *available* quantity at that store; the difference is booked as
  an inventory adjustment through `inventory.services.apply_movement` (reason "Catalogue import"),
  never written directly.
- Imports create real rows (`is_demo` stays False) and refuse to change demo products, demo brands
  and demo stores unless `allow_demo` is set.
- Pictures are not imported: upload them in the dashboard (product › Pictures). The files are never
  downloaded from links in a spreadsheet.
"""

import csv
import datetime
import io
import os
import re
import uuid
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation

from django.core.exceptions import ValidationError as DjangoValidationError
from django.db import IntegrityError, transaction
from django.utils.text import slugify

from apps.core.exceptions import ServiceError

from . import services
from .models import Brand, Category, Product, ProductRelation, ProductVariant, normalize_sku
from .taxonomy import APPROVED_TAXONOMY

MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_ROWS = 5000
MAX_COLUMNS = 80
FILE_TYPES = (".csv", ".xlsx")
OPTION_PREFIX = "option:"
STOCK_REASON = "Catalogue import"

RELATION_COLUMNS = {f"related_{kind}": kind for kind in ProductRelation.Kind.values}

# (column, what it holds), in template and export order.
COLUMNS: list[tuple[str, str]] = [
    ("product_sku", "Product (model) SKU. Rows with the same product_sku are variants of one product. Required."),
    ("product_name", "Product name (up to 200 characters). Required for a new product."),
    (
        "category",
        "One of the ten approved departments, by name or slug, e.g. 'Power Tools & Accessories' or "
        "'power-tools'. Required for a new product.",
    ),
    ("subcategory", "An approved subcategory of that department, by name or slug, e.g. 'Drills'. Blank = none."),
    ("brand", "Brand name. A name that does not exist yet creates the brand. Blank = no brand."),
    ("unit", "piece, pack, box, m, coil, kg, l, bag, pair or set. Blank = unchanged (new products: piece)."),
    ("gst_rate", "GST % included in the price, 0 to 40, e.g. 18. Required for a new product."),
    ("hsn_code", "HSN code, 4 to 8 digits. Format the column as text to keep leading zeros."),
    ("max_order_quantity", "Most units of the product in one order. Blank = unchanged (new products: 50)."),
    ("short_description", "One-line summary (up to 300 characters)."),
    ("description", "Full description."),
    ("tags", "Search keywords separated by commas, e.g. 'pvc, 1/2 inch'."),
    (
        "specifications",
        "Name=value pairs separated by ';', with an optional unit in brackets: 'Thickness=18 [mm]; Grade=BWP'.",
    ),
    ("is_published", "yes or no: visible to customers. Blank = unchanged (new products start unpublished)."),
    *[
        (column, f"Product SKUs linked as '{ProductRelation.Kind(kind).label}', separated by commas.")
        for column, kind in RELATION_COLUMNS.items()
    ],
    ("variant_sku", "Variant SKU (unique across the shop). Required on every row."),
    ("variant_name", "Variant name, e.g. '2.5 sq mm, 90 m coil'. Blank for a simple product."),
    (
        "attributes",
        "Variant options as Name=value pairs separated by ';', e.g. 'Size=2.5 sq mm; Colour=Red'. Extra columns "
        "named 'option:<Name>' add one option each.",
    ),
    ("mrp", "MRP per unit in rupees, GST inclusive. Required for a new variant."),
    ("price", "Selling price per unit, GST inclusive, not above the MRP. Required for a new variant."),
    ("cost_price", "Purchase cost per unit, GST inclusive. Optional; used only for profit reports, never shown."),
    (
        "bulk_tiers",
        "Bulk prices as quantity:price pairs, smallest quantity first, e.g. '10:95; 50:90'. Each price must be "
        "lower than the one before (and the normal price).",
    ),
    (
        "stock",
        "Available stock per store code, e.g. 'BLR-01:25; BLR-02:10'. Sets the available quantity at those "
        "stores (booked as an adjustment). Blank = stock unchanged.",
    ),
]
COLUMN_NAMES = [name for name, _ in COLUMNS]
PRODUCT_COLUMNS = (
    "product_name", "category", "subcategory", "brand", "unit", "gst_rate", "hsn_code", "max_order_quantity",
    "short_description", "description", "tags", "specifications", "is_published", *RELATION_COLUMNS,
)  # fmt: skip
REQUIRED_COLUMNS = ("product_sku", "variant_sku")


class ImportFileError(ServiceError):
    """The file itself can't be read (type, size, encoding, header)."""

    default_code = "invalid_import_file"


class _CellError(Exception):
    pass


# --- Reading the file ------------------------------------------------------------------------------


@dataclass
class Table:
    columns: list[str]
    rows: list[tuple[int, dict[str, str]]]  # (row number in the sheet, {column: text})


def read_table(content: bytes, filename: str, *, sheet: str | None = None) -> Table:
    """The header and the non-blank rows of a .csv or .xlsx file, every cell as trimmed text."""
    extension = os.path.splitext((filename or "").lower())[1]
    if extension not in FILE_TYPES:
        raise ImportFileError("Upload a .csv or .xlsx file.", code="unsupported_file_type")
    if len(content) > MAX_FILE_BYTES:
        raise ImportFileError(
            f"The file is larger than {MAX_FILE_BYTES // (1024 * 1024)} MB. Split it into smaller files.",
            code="file_too_large",
        )
    if not content.strip():
        raise ImportFileError("The file is empty.", code="empty_file")
    grid = _xlsx_grid(content, sheet) if extension == ".xlsx" else _csv_grid(content)

    header = None
    rows: list[tuple[int, dict[str, str]]] = []
    for number, values in grid:
        values = list(values)[:MAX_COLUMNS]
        if header is None:
            header = [_header(v) for v in values]
            _check_header(header)
            continue
        cells = {column: _text(values[i]) if i < len(values) else "" for i, column in enumerate(header) if column}
        if not any(cells.values()):
            continue
        if len(rows) >= MAX_ROWS:
            raise ImportFileError(
                f"The file has more than {MAX_ROWS} rows. Split it into smaller files.", "too_many_rows"
            )
        rows.append((number, cells))
    if header is None:
        raise ImportFileError("The file is empty.", code="empty_file")
    return Table(columns=[c for c in header if c], rows=rows)


def _csv_grid(content: bytes):
    if content.startswith(b"PK\x03\x04"):
        raise ImportFileError("This is an Excel file: name it .xlsx, or save it as CSV.", code="unsupported_file_type")
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        raise ImportFileError(
            "The CSV file is not UTF-8. In Excel choose 'CSV UTF-8 (Comma delimited)' when saving.", "bad_encoding"
        ) from None
    first_line = text.split("\n", 1)[0]
    delimiter = ","
    if "," not in first_line:  # some Excel locales save with ';' or tabs
        delimiter = ";" if ";" in first_line else ("\t" if "\t" in first_line else ",")
    try:
        return list(enumerate(csv.reader(io.StringIO(text), delimiter=delimiter), start=1))
    except csv.Error as exc:
        raise ImportFileError(f"The CSV file could not be read: {exc}.", code="bad_csv") from None


def _xlsx_grid(content: bytes, sheet: str | None):
    if not content.startswith(b"PK\x03\x04"):
        raise ImportFileError("This is not a valid .xlsx file.", code="bad_xlsx")
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception:  # noqa: BLE001 — any parse failure is "not a readable workbook"
        raise ImportFileError("This .xlsx file could not be opened. Save it again in Excel.", "bad_xlsx") from None
    if sheet:
        if sheet not in workbook.sheetnames:
            raise ImportFileError(
                f"There is no sheet '{sheet}'. Sheets: {', '.join(workbook.sheetnames)}.", code="unknown_sheet"
            )
        worksheet = workbook[sheet]
    else:
        worksheet = workbook.worksheets[0]
    rows = []  # (row number, values); blank rows are left out (sheets often have long formatted tails)
    for number, values in enumerate(worksheet.iter_rows(values_only=True, max_col=MAX_COLUMNS), start=1):
        if number == 1 or any(v not in (None, "") for v in values):
            rows.append((number, values))
            if len(rows) > MAX_ROWS + 1:
                break  # read_table reports "too many rows"
    workbook.close()
    return rows


def _header(value) -> str:
    raw = " ".join(str(value or "").split())
    if raw.lower().startswith(OPTION_PREFIX):
        name = raw[len(OPTION_PREFIX) :].strip()
        return OPTION_PREFIX + name if name else ""
    return re.sub(r"[\s\-]+", "_", raw.lower())


def _check_header(header: list[str]) -> None:
    named = [c for c in header if c]
    unknown = [c for c in named if c not in COLUMN_NAMES and not c.startswith(OPTION_PREFIX)]
    if unknown:
        raise ImportFileError(
            f"Unknown column(s): {', '.join(unknown)}. Download the template for the column names.",
            code="unknown_columns",
            details={"unknown_columns": unknown},
        )
    seen = set()
    for column in named:
        key = column.lower()
        if key in seen:
            raise ImportFileError(f"The column '{column}' appears twice.", code="duplicate_column")
        seen.add(key)
    if missing := [c for c in REQUIRED_COLUMNS if c not in named]:
        raise ImportFileError(f"Missing column(s): {', '.join(missing)}.", code="missing_columns")


def _text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        text = "yes" if value else "no"
    elif isinstance(value, float):
        text = str(int(value)) if value.is_integer() else repr(value)
    elif isinstance(value, datetime.datetime | datetime.date):
        text = value.isoformat()
    else:
        text = str(value)
    text = text.replace("\x00", "").strip()
    if text[:1] == "'" and text[1:2] in ("=", "+", "-", "@"):
        text = text[1:]  # the export's protection against spreadsheet formulas
    return text


# --- Cell parsers ----------------------------------------------------------------------------------


def _decimal(text: str, *, places: int = 2) -> Decimal:
    cleaned = text.replace("₹", "").replace(",", "").replace("%", "").strip()
    try:
        value = Decimal(cleaned)
    except InvalidOperation:
        raise _CellError(f"'{text}' is not a number.") from None
    if not value.is_finite():
        raise _CellError(f"'{text}' is not a number.")
    if value.as_tuple().exponent < -places:
        raise _CellError(f"'{text}' has more than {places} decimal places.")
    return value.quantize(Decimal(1).scaleb(-places))


def _money(text: str, *, allow_zero: bool = False) -> Decimal:
    value = _decimal(text)
    if value < 0 or (value == 0 and not allow_zero):
        raise _CellError("Enter an amount more than 0." if not allow_zero else "The amount can't be negative.")
    if value >= Decimal("10000000000"):
        raise _CellError("The amount is too large.")
    return value


def _whole(text: str, *, minimum: int, maximum: int) -> int:
    value = _decimal(text, places=0) if text else None
    if value is None or value < minimum or value > maximum:
        raise _CellError(f"Enter a whole number from {minimum} to {maximum}.")
    return int(value)


def _bool(text: str) -> bool:
    lowered = text.lower()
    if lowered in ("yes", "y", "true", "1"):
        return True
    if lowered in ("no", "n", "false", "0"):
        return False
    raise _CellError("Enter yes or no.")


def _pairs(text: str, separator: str = "=") -> list[tuple[str, str]]:
    """'A=1; B=2' → [('A', '1'), ('B', '2')]."""
    pairs = []
    for part in text.split(";"):
        part = part.strip()
        if not part:
            continue
        if separator not in part:
            raise _CellError(f"'{part}' should look like name{separator}value.")
        key, value = (s.strip() for s in part.split(separator, 1))
        if not key or not value:
            raise _CellError(f"'{part}' needs both a name and a value.")
        pairs.append((key, value))
    return pairs


def _length(text: str, limit: int) -> str:
    if len(text) > limit:
        raise _CellError(f"Too long: at most {limit} characters.")
    return text


def _sku(text: str) -> str:
    sku = normalize_sku(text)
    if len(sku) > 64:
        raise _CellError("A SKU can have at most 64 characters.")
    return sku


def _sku_list(text: str) -> list[str]:
    skus = []
    for part in re.split(r"[,;]", text):
        if sku := _sku(part):
            if sku not in skus:
                skus.append(sku)
    return skus


def parse_specifications(text: str) -> list[dict]:
    specs, seen = [], set()
    for name, value in _pairs(text):
        unit = ""
        if match := re.fullmatch(r"(.*?)\s*\[([^\]]*)\]", value):
            value, unit = match.group(1).strip(), match.group(2).strip()
        if not value:
            raise _CellError(f"'{name}' needs a value.")
        if len(name) > 80 or len(value) > 200 or len(unit) > 20:
            raise _CellError(f"'{name}' is too long (name 80, value 200, unit 20 characters).")
        key = slugify(name)
        if key in seen:
            raise _CellError(f"'{name}' is listed twice.")
        seen.add(key)
        specs.append({"name": name, "value": value, "unit": unit})
    return specs


def parse_tiers(text: str) -> list[tuple[int, Decimal]]:
    tiers = []
    for part in text.split(";"):
        part = part.strip()
        if not part:
            continue
        if ":" not in part:
            raise _CellError(f"'{part}' should look like quantity:price, e.g. 10:95.")
        quantity, price = (s.strip() for s in part.split(":", 1))
        tiers.append((_whole(quantity, minimum=2, maximum=100000), _money(price)))
    quantities = [q for q, _ in tiers]
    if quantities != sorted(set(quantities)):
        raise _CellError("List each quantity once, smallest first, e.g. '10:95; 50:90'.")
    return tiers


def parse_stock(text: str) -> list[tuple[str, int]]:
    rows, seen = [], set()
    for code, quantity in _pairs(text, ":"):
        code = code.strip().upper()
        if code in seen:
            raise _CellError(f"Store {code} is listed twice.")
        seen.add(code)
        rows.append((code, _whole(quantity, minimum=0, maximum=10_000_000)))
    return rows


# --- Planning (validation) -------------------------------------------------------------------------


@dataclass
class VariantPlan:
    row: int
    sku: str
    existing: ProductVariant | None
    fields: dict = field(default_factory=dict)
    tiers: list[tuple[int, Decimal]] | None = None
    stock: list[tuple[object, int]] | None = None  # [(Store, wanted available quantity)]


@dataclass
class ProductPlan:
    sku: str
    rows: list[int]
    existing: Product | None
    fields: dict = field(default_factory=dict)
    specifications: list[dict] | None = None
    relations: dict[str, list[str]] = field(default_factory=dict)
    variants: list[VariantPlan] = field(default_factory=list)


@dataclass
class ImportReport:
    dry_run: bool
    filename: str = ""
    committed: bool = False
    rows: list[dict] = field(default_factory=list)  # {row, product_sku, variant_sku, action}
    errors: list[dict] = field(default_factory=list)  # {row, column, sku, message}
    counts: dict = field(default_factory=dict)

    @property
    def ok(self) -> bool:
        return not self.errors

    def error(self, row: int | None, message: str, *, column: str = "", sku: str = "") -> None:
        self.errors.append({"row": row, "column": column, "sku": sku, "message": message})

    def as_dict(self) -> dict:
        return {
            "ok": self.ok,
            "dry_run": self.dry_run,
            "committed": self.committed,
            "filename": self.filename,
            "counts": self.counts,
            "rows": self.rows,
            "errors": sorted(self.errors, key=lambda e: (e["row"] or 0, e["column"])),
        }


def _department_index() -> dict[str, tuple[str, dict[str, str]]]:
    """name/slug (case-folded) → (department slug, {subcategory name/slug → subcategory slug})."""
    index = {}
    for (slug, name), children in APPROVED_TAXONOMY.items():
        subs = {}
        for child_slug, child_name in children:
            subs[child_slug.casefold()] = subs[child_name.casefold()] = child_slug
        index[slug.casefold()] = index[name.casefold()] = (slug, subs)
    return index


class _Planner:
    def __init__(self, table: Table, report: ImportReport, *, allow_demo: bool):
        self.table, self.report, self.allow_demo = table, report, allow_demo
        self.columns = set(table.columns)
        self.option_columns = [c for c in table.columns if c.startswith(OPTION_PREFIX)]

    def err(self, row: int, column: str, message: str, sku: str = "") -> None:
        self.report.error(row, message, column=column, sku=sku)

    def plan(self) -> list[ProductPlan]:
        from apps.stores.models import Store

        groups: dict[str, list[tuple[int, dict]]] = {}
        variant_rows: dict[str, int] = {}
        for number, cells in self.table.rows:
            try:
                product_sku = _sku(cells.get("product_sku", ""))
            except _CellError as exc:
                self.err(number, "product_sku", str(exc))
                continue
            try:
                variant_sku = _sku(cells.get("variant_sku", ""))
            except _CellError as exc:
                self.err(number, "variant_sku", str(exc), product_sku)
                continue
            if not product_sku:
                self.err(number, "product_sku", "Enter the product SKU.", variant_sku)
                continue
            if not variant_sku:
                self.err(number, "variant_sku", "Enter the variant SKU.", product_sku)
                continue
            if variant_sku in variant_rows:
                self.err(number, "variant_sku", f"Duplicate row: SKU already on row {variant_rows[variant_sku]}.")
                continue
            variant_rows[variant_sku] = number
            groups.setdefault(product_sku, []).append((number, cells))

        self.products = {p.sku: p for p in Product.all_objects.select_related("category").filter(sku__in=groups)}
        self.variants = {
            v.sku: v for v in ProductVariant.objects.select_related("product").filter(sku__in=variant_rows)
        }
        self.brands = {b.name.casefold(): b for b in Brand.all_objects.all()}
        self.stores = {s.code: s for s in Store.all_objects.all()}
        self.categories = {c.slug: c for c in Category.objects.filter(parent__isnull=True)}
        self.subcategories = {c.slug: c for c in Category.objects.filter(parent__isnull=False)}
        self.departments = _department_index()
        self.file_products = set(groups)
        related = set()
        for rows in groups.values():
            for _, cells in rows:
                for column in RELATION_COLUMNS:
                    try:
                        related.update(_sku_list(cells.get(column, "")))
                    except _CellError:
                        pass
        self.related = {p.sku: p for p in Product.all_objects.filter(sku__in=related)}
        return [plan for sku, rows in groups.items() if (plan := self._product(sku, rows))]

    # One product -------------------------------------------------------------------------------

    def _product(self, sku: str, rows: list[tuple[int, dict]]) -> ProductPlan | None:
        first = rows[0][0]
        existing = self.products.get(sku)
        before = len(self.report.errors)
        if existing is not None and existing.is_deleted:
            self.err(first, "product_sku", f"Product {sku} is deleted. Restore it in the dashboard first.", sku)
            return None
        if existing is not None and existing.is_demo and not self.allow_demo:
            self.err(first, "product_sku", f"Product {sku} is demo data and is not changed by imports.", sku)
            return None
        plan = ProductPlan(sku=sku, rows=[n for n, _ in rows], existing=existing)
        values = self._product_values(rows, sku)
        self._product_fields(plan, values, first)
        for number, cells in rows:
            if variant := self._variant(plan, number, cells):
                plan.variants.append(variant)
        return plan if len(self.report.errors) == before else None

    def _product_values(self, rows, sku) -> dict[str, str]:
        """Each product column's value: the first filled cell; a different filled value later is an error."""
        values = {}
        for column in PRODUCT_COLUMNS:
            if column not in self.columns:
                continue
            value, origin = "", None
            for number, cells in rows:
                text = cells.get(column, "")
                if not text:
                    continue
                if origin is None:
                    value, origin = text, number
                elif text != value:
                    self.err(number, column, f"Differs from row {origin} of the same product ({sku}).", sku)
            values[column] = value
        return values

    def _product_fields(self, plan: ProductPlan, values: dict, row: int) -> None:
        existing, fields, sku = plan.existing, plan.fields, plan.sku

        def parse(column, parser):
            if column not in values:
                return
            try:
                parser(values[column])
            except _CellError as exc:
                self.err(row, column, str(exc), sku)

        def name(text):
            if not text:
                raise _CellError("Enter the product name.")
            fields["name"] = _length(text, 200)

        def unit(text):
            if not text:
                return
            for choice, label in Product.Unit.choices:
                if text.lower() in (choice, label.lower()):
                    fields["unit"] = choice
                    return
            raise _CellError(f"Unknown unit '{text}'. Use one of: {', '.join(Product.Unit.values)}.")

        def gst(text):
            if not text:
                raise _CellError("Enter the GST rate (0 to 40).")
            rate = _decimal(text)
            if not Decimal(0) <= rate <= Decimal(40):
                raise _CellError("GST must be from 0 to 40.")
            fields["gst_rate"] = rate

        def hsn(text):
            if text and not re.fullmatch(r"\d{4,8}", text):
                raise _CellError("An HSN code is 4 to 8 digits.")
            fields["hsn_code"] = text

        def max_quantity(text):
            if text:
                fields["max_order_quantity"] = _whole(text, minimum=1, maximum=32767)

        def tags(text):
            items = []
            for tag in text.split(","):
                tag = tag.strip()
                if tag and tag not in items:
                    items.append(_length(tag, 50))
            if len(items) > 40:
                raise _CellError("At most 40 tags.")
            fields["tags"] = items

        def specifications(text):
            plan.specifications = parse_specifications(text)

        def published(text):
            if text:
                fields["is_published"] = _bool(text)

        def brand(text):
            if not text:
                fields["brand"] = None
                return
            found = self.brands.get(text.casefold())
            if found is None:
                fields["brand"] = _length(text, 120)  # created when the import runs
            elif found.is_deleted:
                raise _CellError(f"The brand '{found.name}' is deleted. Restore it or use another brand.")
            elif found.is_demo and not self.allow_demo:
                raise _CellError(f"The brand '{found.name}' is demo data; use a real brand.")
            else:
                fields["brand"] = found

        parse("product_name", name)
        parse("unit", unit)
        parse("gst_rate", gst)
        parse("hsn_code", hsn)
        parse("max_order_quantity", max_quantity)
        parse("short_description", lambda t: fields.__setitem__("short_description", _length(t, 300)))
        parse("description", lambda t: fields.__setitem__("description", _length(t, 20000)))
        parse("tags", tags)
        parse("specifications", specifications)
        parse("is_published", published)
        parse("brand", brand)
        self._category(plan, values, row)
        for column, kind in RELATION_COLUMNS.items():
            parse(column, lambda t, kind=kind: plan.relations.__setitem__(kind, self._related(t, sku)))
        if existing is None:
            for column, label in (("product_name", "product name"), ("category", "category"), ("gst_rate", "GST")):
                if column not in values:
                    self.err(row, column, f"A new product needs its {label}.", sku)

    def _category(self, plan: ProductPlan, values: dict, row: int) -> None:
        sku, fields, existing = plan.sku, plan.fields, plan.existing
        department = None
        if "category" in values:
            text = values["category"]
            found = self.departments.get(text.casefold()) if text else None
            if not text:
                self.err(row, "category", "Enter the category.", sku)
                return
            if found is None:
                self.err(row, "category", f"'{text}' is not one of the approved categories.", sku)
                return
            department = found
            category = self.categories.get(found[0])
            if category is None:
                self.err(row, "category", "The approved categories are missing: run manage.py apply_taxonomy.", sku)
                return
            fields["category"] = category
            moved = existing is not None and existing.subcategory_id and existing.subcategory.parent_id != category.pk
            if moved and "subcategory" not in values:
                self.err(
                    row, "category", "The product's subcategory is in another category: add a subcategory column.", sku
                )
                return
        elif existing is not None:
            department = self.departments.get(existing.category.slug.casefold())
        if "subcategory" not in values:
            return
        text = values["subcategory"]
        if not text:
            fields["subcategory"] = None
            return
        if department is None:
            self.err(row, "subcategory", "Give the approved category too.", sku)
            return
        slug = department[1].get(text.casefold())
        if slug is None or slug not in self.subcategories:
            self.err(row, "subcategory", f"'{text}' is not a subcategory of the chosen category.", sku)
            return
        fields["subcategory"] = self.subcategories[slug]
        if "category" not in fields and existing is not None and existing.category.slug != department[0]:
            fields["category"] = self.categories[department[0]]

    def _related(self, text: str, sku: str) -> list[str]:
        skus = _sku_list(text)
        for other in skus:
            if other == sku:
                raise _CellError("A product can't be linked to itself.")
            found = self.related.get(other)
            if other in self.file_products:
                continue
            if found is None or found.is_deleted:
                raise _CellError(f"No product with SKU {other}.")
            if found.is_demo and not self.allow_demo:
                raise _CellError(f"{other} is demo data; real products can't link to it.")
        return skus

    # One variant -------------------------------------------------------------------------------

    def _variant(self, plan: ProductPlan, row: int, cells: dict) -> VariantPlan | None:
        sku = _sku(cells["variant_sku"])
        existing = self.variants.get(sku)
        if existing is not None and existing.product.sku != plan.sku:
            self.err(row, "variant_sku", f"SKU {sku} already belongs to product {existing.product.sku}.", sku)
            return None
        variant = VariantPlan(row=row, sku=sku, existing=existing)
        fields = variant.fields
        before = len(self.report.errors)

        def parse(column, parser):
            if column not in self.columns:
                return
            try:
                parser(cells.get(column, ""))
            except _CellError as exc:
                self.err(row, column, str(exc), sku)

        def money(column, *, required):
            def inner(text):
                if not text:
                    if required:
                        raise _CellError(f"Enter the {column.upper() if column == 'mrp' else column}.")
                    fields[column] = None
                    return
                fields[column] = _money(text, allow_zero=column == "cost_price")

            return inner

        parse("variant_name", lambda t: fields.__setitem__("name", _length(t, 120)))
        parse("mrp", money("mrp", required=True))
        parse("price", money("price", required=True))
        parse("cost_price", money("cost_price", required=False))
        self._attributes(variant, cells)
        if existing is None:
            for column in ("mrp", "price"):
                if column not in self.columns:
                    self.err(row, column, f"A new variant needs its {column.upper() if column == 'mrp' else column}.")
        mrp = fields.get("mrp", getattr(existing, "mrp", None))
        price = fields.get("price", getattr(existing, "price", None))
        if mrp and price and price > mrp:
            self.err(row, "price", f"The price ₹{price} is above the MRP ₹{mrp}.", sku)
        parse("bulk_tiers", lambda t: setattr(variant, "tiers", parse_tiers(t)))
        if variant.tiers and mrp and price and price <= mrp:
            limit = plan.fields.get("max_order_quantity") or getattr(plan.existing, "max_order_quantity", None) or 50
            probe = ProductVariant(product=Product(max_order_quantity=limit), price=price, mrp=mrp)
            try:
                services.validate_price_tiers(probe, [{"min_quantity": q, "unit_price": p} for q, p in variant.tiers])
            except services.CatalogRuleError as exc:
                self.err(row, "bulk_tiers", str(exc.detail), sku)
        if cells.get("stock"):
            parse("stock", lambda t: self._stock(variant, t))
        return variant if len(self.report.errors) == before else None

    def _attributes(self, variant: VariantPlan, cells: dict) -> None:
        if "attributes" not in self.columns and not self.option_columns:
            return
        attributes: dict[str, str] = {}
        try:
            for key, value in _pairs(cells.get("attributes", "")):
                if key.casefold() in (k.casefold() for k in attributes):
                    raise _CellError(f"'{key}' is listed twice.")
                attributes[_length(key, 60)] = _length(value, 80)
        except _CellError as exc:
            self.err(variant.row, "attributes", str(exc), variant.sku)
            return
        for column in self.option_columns:
            name, value = column[len(OPTION_PREFIX) :], cells.get(column, "")
            if not value:
                continue
            same = next((k for k in attributes if k.casefold() == name.casefold()), None)
            if same is not None and attributes[same] != value:
                self.err(variant.row, column, f"Differs from '{same}' in the attributes column.", variant.sku)
                return
            attributes[same or _length(name, 60)] = value[:80]
        variant.fields["attributes"] = attributes

    def _stock(self, variant: VariantPlan, text: str) -> None:
        rows = []
        for code, quantity in parse_stock(text):
            store = self.stores.get(code)
            if store is None or store.is_deleted:
                raise _CellError(f"No store with code {code}.")
            if store.is_demo and not self.allow_demo:
                raise _CellError(f"Store {code} is a demo store; real stock can't be put there.")
            rows.append((store, quantity))
        variant.stock = rows


# --- Running the import ----------------------------------------------------------------------------


def _error_text(exc) -> str:
    if isinstance(exc, DjangoValidationError):
        if hasattr(exc, "error_dict"):
            return "; ".join(f"{k}: {' '.join(v)}" for k, v in exc.message_dict.items())
        return " ".join(exc.messages)
    if isinstance(exc, ServiceError):
        return str(exc.detail)
    return "The row could not be saved (it conflicts with existing data)."


class _Runner:
    def __init__(self, report: ImportReport, *, actor, batch: str):
        self.report, self.actor, self.batch = report, actor, batch
        self.new_brands: dict[str, Brand] = {}
        self.counts = {"stock_movements": 0, "brands_created": 0}

    def brand(self, name: str) -> Brand:
        key = name.casefold()
        if key not in self.new_brands:
            brand = Brand(name=name)
            brand.save()
            self.new_brands[key] = brand
            self.counts["brands_created"] += 1
        return self.new_brands[key]

    def product(self, plan: ProductPlan) -> Product:
        data = dict(plan.fields)
        if isinstance(data.get("brand"), str):
            data["brand"] = self.brand(data["brand"])
        if plan.existing is None:
            data["sku"] = plan.sku
        if plan.specifications is not None:
            data["specifications"] = plan.specifications
        payloads = []
        for variant in plan.variants:
            payload = {"sku": variant.sku, **variant.fields}
            if variant.existing is not None:
                payload["id"] = variant.existing.pk
            if variant.tiers is not None:
                payload["price_tiers"] = [{"min_quantity": q, "unit_price": p} for q, p in variant.tiers]
            payloads.append(payload)
        data["variants"] = payloads
        product = services.save_product(instance=plan.existing, data=data, actor=self.actor)
        self._options(product, plan)
        self._stock(plan)
        return product

    def _options(self, product: Product, plan: ProductPlan) -> None:
        """Attributes whose name is an option type of the product's category also become structured
        options (filters, variant pickers); the rest stay display attributes."""
        option_types = {t.name.casefold(): t for t in services.option_types_for(product)}
        changed = False
        for plan_variant in plan.variants:
            if "attributes" not in plan_variant.fields:
                continue
            variant = ProductVariant.objects.get(sku=plan_variant.sku)
            attributes = {
                (option_types[k.casefold()].name if k.casefold() in option_types else k): v
                for k, v in plan_variant.fields["attributes"].items()
            }
            if attributes != variant.attributes:
                ProductVariant.objects.filter(pk=variant.pk).update(attributes=attributes)
                variant.attributes = attributes
            options = [
                {"option_type": option_types[k.casefold()].pk, "value": v}
                for k, v in attributes.items()
                if k.casefold() in option_types
            ]
            if options or variant.option_values.exists():
                services.set_variant_options(variant, options)
                changed = True
        if changed:
            services._check_product_rules(product)  # two variants can't end up with the same options

    def _stock(self, plan: ProductPlan) -> None:
        from apps.inventory.models import Inventory, StockMovement
        from apps.inventory.services import apply_movement

        for plan_variant in plan.variants:
            if not plan_variant.stock:
                continue
            variant = ProductVariant.objects.get(sku=plan_variant.sku)
            for store, wanted in plan_variant.stock:
                inventory = Inventory.objects.filter(store=store, variant=variant).first()
                if inventory is None:
                    if wanted == 0:
                        continue
                    inventory = Inventory.objects.create(store=store, variant=variant)
                change = wanted - inventory.available_quantity
                if change:
                    apply_movement(
                        inventory,
                        StockMovement.Type.ADJUSTMENT,
                        change,
                        reason=STOCK_REASON,
                        reference_type="catalog_import",
                        reference_id=self.batch,
                        actor=self.actor,
                    )
                    self.counts["stock_movements"] += 1

    def relations(self, plans: list[ProductPlan]) -> None:
        """Replace each listed link type with the products in the file (keeping notes of links that stay)."""
        wanted_skus = {sku for plan in plans for skus in plan.relations.values() for sku in skus}
        by_sku = dict(Product.objects.filter(sku__in=wanted_skus).values_list("sku", "pk"))
        for plan in plans:
            if not plan.relations:
                continue
            product = Product.all_objects.get(sku=plan.sku)
            for kind, skus in plan.relations.items():
                ids = [by_sku[s] for s in skus]
                current = product.relations.filter(kind=kind)
                current.exclude(related_id__in=ids).delete()
                have = set(current.values_list("related_id", flat=True))
                for order, related_id in enumerate(ids):
                    if related_id in have:
                        ProductRelation.objects.filter(product=product, related_id=related_id, kind=kind).update(
                            sort_order=order
                        )
                    else:
                        ProductRelation.objects.create(
                            product=product, related_id=related_id, kind=kind, sort_order=order
                        )


def run_import(
    table: Table, *, dry_run: bool = True, allow_demo: bool = False, actor=None, filename: str = ""
) -> ImportReport:
    """Validate every row, then import all of it in one transaction (rolled back for a dry run or on
    any error). The report lists each row's action (create, update, unchanged) and all errors."""
    from .importer_export import snapshot

    report = ImportReport(dry_run=dry_run, filename=filename)
    plans = _Planner(table, report, allow_demo=allow_demo).plan()
    report.counts = {"rows": len(table.rows)}
    if not table.rows:
        report.error(None, "The file has no product rows.")
    if not report.ok:
        return report

    existing_ids = [v.existing.pk for plan in plans for v in plan.variants if v.existing is not None]
    before = snapshot(existing_ids)
    runner = _Runner(report, actor=actor, batch=uuid.uuid4().hex[:12])
    with transaction.atomic():
        for plan in plans:
            try:
                with transaction.atomic():
                    runner.product(plan)
            except (ServiceError, DjangoValidationError, IntegrityError) as exc:
                report.error(plan.rows[0], _error_text(exc), sku=plan.sku)
        if report.ok:
            try:
                with transaction.atomic():
                    runner.relations(plans)
            except (ServiceError, DjangoValidationError, IntegrityError) as exc:
                report.error(None, _error_text(exc))
        if report.ok:
            variant_ids = dict(
                ProductVariant.objects.filter(sku__in=[v.sku for p in plans for v in p.variants]).values_list(
                    "sku", "pk"
                )
            )
            after = snapshot(list(variant_ids.values()))
            counts = {"products_created": 0, "products_updated": 0, "variants_created": 0, "variants_updated": 0}
            counts["unchanged"] = 0
            for plan in plans:
                product_changed = False
                for variant in plan.variants:
                    if variant.existing is None:
                        action = "create"
                        counts["variants_created"] += 1
                    elif before.get(variant.existing.pk) == after.get(variant_ids[variant.sku]):
                        action = "unchanged"
                        counts["unchanged"] += 1
                    else:
                        action = "update"
                        counts["variants_updated"] += 1
                    product_changed |= action != "unchanged"
                    report.rows.append(
                        {"row": variant.row, "product_sku": plan.sku, "variant_sku": variant.sku, "action": action}
                    )
                if plan.existing is None:
                    counts["products_created"] += 1
                elif product_changed:
                    counts["products_updated"] += 1
            report.counts.update(counts, **runner.counts)
        report.committed = report.ok and not dry_run
        if not report.committed:
            transaction.set_rollback(True)
    report.rows.sort(key=lambda r: r["row"])
    if report.committed:
        from apps.core.audit import record_audit

        record_audit(action="catalog.import", actor=actor, changes={"file": filename, "counts": report.counts})
    return report


def import_file(content: bytes, filename: str, *, sheet: str | None = None, **options) -> ImportReport:
    return run_import(read_table(content, filename, sheet=sheet), filename=filename, **options)
