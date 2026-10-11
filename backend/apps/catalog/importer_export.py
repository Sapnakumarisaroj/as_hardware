"""Catalogue export in the import's own columns (CSV or .xlsx), and the blank template.

An export can be edited and imported again unchanged: same columns, one row per variant, product
columns repeated on every variant row. Deleted products and demo data (`is_demo`) are left out
unless asked for. Text that a spreadsheet could run as a formula (=, +, -, @) is prefixed with ',
which the importer strips again.
"""

import csv
import io
from collections import defaultdict

from django.db.models import Prefetch

from .importer import COLUMNS, RELATION_COLUMNS
from .models import Product, ProductRelation, ProductVariant

HEADER = [name for name, _ in COLUMNS]
NUMERIC_COLUMNS = ("gst_rate", "max_order_quantity", "mrp", "price", "cost_price")


def _department_names(product: Product) -> tuple[str, str]:
    """(category, subcategory) names as the importer reads them: department + approved subcategory."""
    category, subcategory = product.category, product.subcategory
    if subcategory is None and category.parent_id is not None:  # an older product filed directly under a sub
        return category.parent.name, category.name
    return category.name, subcategory.name if subcategory else ""


def _spec_text(spec) -> str:
    return f"{spec.name}={spec.value}" + (f" [{spec.unit}]" if spec.unit else "")


def _product_columns(product: Product) -> dict:
    category, subcategory = _department_names(product)
    relations = defaultdict(list)
    for relation in product.relations.all():
        if not relation.related.is_deleted:
            relations[f"related_{relation.kind}"].append(relation.related.sku)
    return {
        "product_sku": product.sku,
        "product_name": product.name,
        "category": category,
        "subcategory": subcategory,
        "brand": product.brand.name if product.brand_id else "",
        "unit": product.unit,
        "gst_rate": product.gst_rate,
        "hsn_code": product.hsn_code,
        "max_order_quantity": product.max_order_quantity,
        "short_description": product.short_description,
        "description": product.description,
        "tags": ", ".join(product.tags or []),
        "specifications": "; ".join(_spec_text(s) for s in product.specifications.all()),
        "highlights": product.highlights,
        "uses": product.uses,
        "installation_steps": product.installation_steps,
        "is_published": "yes" if product.is_published else "no",
        **{column: ", ".join(relations.get(column, [])) for column in RELATION_COLUMNS},
    }


def _variant_columns(variant: ProductVariant, stock: dict[int, list[tuple[str, int]]]) -> dict:
    return {
        "variant_sku": variant.sku,
        "variant_name": variant.name,
        "attributes": "; ".join(f"{k}={v}" for k, v in (variant.attributes or {}).items()),
        "mrp": variant.mrp,
        "price": variant.price,
        "cost_price": variant.cost_price if variant.cost_price is not None else "",
        "price_approved": "yes" if variant.price_approved else "no",
        "price_source": variant.price_source,
        "bulk_tiers": "; ".join(f"{t.min_quantity}:{t.unit_price}" for t in variant.price_tiers.all()),
        "stock": "; ".join(f"{code}:{quantity}" for code, quantity in stock.get(variant.pk, [])),
    }


def _products(product_filter: dict, variant_ids: list[int] | None = None):
    variants = ProductVariant.objects.prefetch_related("price_tiers").order_by("sort_order", "id")
    if variant_ids is not None:
        variants = variants.filter(pk__in=variant_ids)
    return (
        Product.all_objects.filter(**product_filter)
        .select_related("brand", "category__parent", "subcategory")
        .prefetch_related(
            "specifications",
            Prefetch("relations", queryset=ProductRelation.objects.select_related("related").order_by("sort_order")),
            Prefetch("variants", queryset=variants),
        )
        .order_by("sku")
    )


def _stock_map(variant_ids, *, include_demo: bool) -> dict[int, list[tuple[str, int]]]:
    from apps.inventory.models import Inventory

    rows = Inventory.objects.filter(variant_id__in=variant_ids, is_active=True, store__is_deleted=False)
    if not include_demo:
        rows = rows.filter(store__is_demo=False)
    stock = defaultdict(list)
    for variant_id, code, quantity in rows.order_by("store__code").values_list(
        "variant_id", "store__code", "available_quantity"
    ):
        stock[variant_id].append((code, quantity))
    return stock


def export_rows(*, include_demo: bool = False, include_deleted: bool = False) -> list[dict]:
    """One dict per variant, keyed by the import columns."""
    product_filter = {} if include_deleted else {"is_deleted": False}
    if not include_demo:
        product_filter["is_demo"] = False
    products = list(_products(product_filter))
    stock = _stock_map([v.pk for p in products for v in p.variants.all()], include_demo=include_demo)
    rows = []
    for product in products:
        shared = _product_columns(product)
        rows.extend({**shared, **_variant_columns(variant, stock)} for variant in product.variants.all())
    return rows


def snapshot(variant_ids: list[int]) -> dict[int, dict]:
    """variant id → its export row (the importer compares these to tell 'update' from 'unchanged')."""
    if not variant_ids:
        return {}
    products = _products({"variants__pk__in": variant_ids}, variant_ids).distinct()
    stock = _stock_map(variant_ids, include_demo=True)
    result = {}
    for product in products:
        shared = _product_columns(product)
        for variant in product.variants.all():
            result[variant.pk] = {**shared, **_variant_columns(variant, stock)}
    return result


def _cell(value) -> str:
    from apps.analytics.views import _cell as safe_cell  # the CSV export's formula-injection guard

    return safe_cell(value)


def write_csv(rows: list[dict]) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(HEADER)
    for row in rows:
        writer.writerow([_cell(row.get(column, "")) for column in HEADER])
    return ("﻿" + buffer.getvalue()).encode("utf-8")  # BOM: Excel opens it as UTF-8


def write_xlsx(rows: list[dict], *, with_help: bool = False) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Catalogue"
    sheet.append(HEADER)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    sheet.freeze_panes = "A2"
    for row in rows:
        values = []
        for column in HEADER:
            value = row.get(column, "")
            # Numbers stay numbers; text is protected from formula injection and stored as text.
            values.append(value if column in NUMERIC_COLUMNS and value != "" else _cell(value))
        sheet.append(values)
    for index, column in enumerate(HEADER, start=1):
        sheet.column_dimensions[sheet.cell(row=1, column=index).column_letter].width = max(12, len(column) + 2)
        if column == "hsn_code":
            for cell in sheet.iter_cols(min_col=index, max_col=index, min_row=1):
                for c in cell:
                    c.number_format = "@"
    if with_help:
        help_sheet = workbook.create_sheet("Columns")
        help_sheet.append(["Column", "What to enter"])
        for cell in help_sheet[1]:
            cell.font = Font(bold=True)
        for name, description in COLUMNS:
            help_sheet.append([name, description])
        help_sheet.column_dimensions["A"].width = 22
        help_sheet.column_dimensions["B"].width = 110
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def template(fmt: str) -> bytes:
    """A blank import file: the header row (and, for Excel, a 'Columns' sheet explaining each column)."""
    return write_xlsx([], with_help=True) if fmt == "xlsx" else write_csv([])


CONTENT_TYPES = {
    "csv": "text/csv; charset=utf-8",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}
